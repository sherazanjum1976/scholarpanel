"""Text extraction, deterministic (non-LLM) analysis, scoring and report building."""
import hashlib
import io
import json
import re

import fitz  # PyMuPDF
from docx import Document

from config import (EXCERPT_CHARS, MAX_DOC_CHARS, MAX_UPLOAD_MB, MIN_DOC_CHARS,
                    WEIGHTS)

# key, label, heading regex, required_for_synopsis, required_for_thesis
SECTIONS = [
    ("abstract", "Abstract / Summary", r"abstract|summary|executive summary", True, True),
    ("introduction", "Introduction / Background", r"introduction|background", True, True),
    ("problem", "Problem Statement", r"problem statement|statement of (?:the )?problem|research problem", True, False),
    ("literature", "Literature Review", r"literature review|review of (?:the )?literature|related work|theoretical (?:framework|background)", True, True),
    ("gap", "Research Gap / Rationale", r"research gap|gap in (?:the )?literature|rationale|justification", True, False),
    ("questions", "Research Questions / Hypotheses", r"research questions?|hypothes[ie]s|research hypothes[ie]s", True, True),
    ("objectives", "Objectives", r"(?:research |study )?objectives?|aims?(?: and objectives)?", True, True),
    ("method", "Methodology", r"(?:research )?(?:methodology|methods?|design)|materials and methods", True, True),
    ("significance", "Significance / Expected Outcomes", r"significance|expected (?:outcomes|results|contributions?)|contributions?", True, False),
    ("results", "Results / Findings", r"results?|findings|data analysis", False, True),
    ("discussion", "Discussion", r"discussion", False, True),
    ("conclusion", "Conclusion / Future Work", r"conclusions?|future work|recommendations", False, True),
    ("ethics", "Ethical Considerations", r"ethic(?:s|al)?(?: considerations| approval| statement)?|informed consent", True, True),
    ("timeline", "Timeline / Work Plan", r"timeline|work plan|time ?frame|schedule|gantt", True, False),
    ("references", "References", r"references|bibliography|works cited", True, True),
]


def _hrx(p):
    return re.compile(
        rf"^\s*(?:chapter\s+\w+\s*[:.\-\u2013]*\s*)?(?:\d+(?:\.\d+)*[.):\-\u2013\s]*)?(?:{p})\b", re.I)


HEAD_RES = {k: _hrx(p) for k, _, p, _, _ in SECTIONS}


# ---------------------------------------------------------------- extraction
def clean_text(t: str) -> str:
    t = t.replace("\r", "\n").replace("\x00", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def extract_text(filename: str, data: bytes):
    """Return (text, truncated_flag). Raises ValueError with a user-friendly message."""
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise ValueError(f"File is larger than {MAX_UPLOAD_MB} MB. Please upload a smaller file.")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        try:
            doc = fitz.open(stream=data, filetype="pdf")
            raw = "\n".join(page.get_text() for page in doc)
        except Exception:
            raise ValueError("This PDF could not be read. It may be corrupted or password-protected.")
    elif ext == "docx":
        try:
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for tb in d.tables:
                for row in tb.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            raw = "\n".join(parts)
        except Exception:
            raise ValueError("This DOCX could not be read. Please re-save it as .docx or PDF.")
    elif ext in ("txt", "md"):
        raw = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload PDF, DOCX or TXT.")
    return finalize_text(raw, scanned_hint=(ext == "pdf"))


def finalize_text(raw: str, scanned_hint: bool = False):
    text = clean_text(raw)
    if len(text) < MIN_DOC_CHARS:
        hint = " If this is a scanned PDF, it needs OCR first." if scanned_hint else ""
        raise ValueError(f"Not enough text found (minimum {MIN_DOC_CHARS} characters).{hint}")
    truncated = len(text) > MAX_DOC_CHARS
    return text[:MAX_DOC_CHARS], truncated


# ------------------------------------------------------------ deterministic analysis
def word_count(text: str) -> int:
    return len(re.findall(r"\w+", text))


def detect_doc_type(text: str) -> str:
    head = text[:2500].lower()
    if "synopsis" in head or "proposal" in head:
        return "synopsis"
    n_chapter = len(re.findall(r"^\s*chapter\s+\w+", text, re.I | re.M))
    thesis_marks = sum(m in head for m in ("dedication", "acknowledg", "table of contents",
                                           "list of figures", "declaration", "thesis"))
    if n_chapter >= 3 or thesis_marks >= 2 or word_count(text) > 9000:
        return "thesis"
    return "synopsis"


def section_bodies(text: str):
    lines = text.split("\n")
    found = {}
    for i, line in enumerate(lines):
        s = line.strip()
        if not s or len(s) > 90 or s.endswith((".", ",", ";")):
            continue
        for key, rx in HEAD_RES.items():
            if key not in found and rx.match(s):
                found[key] = i
    order = sorted(found.items(), key=lambda kv: kv[1])
    bodies = {}
    for n, (k, i) in enumerate(order):
        end = order[n + 1][1] if n + 1 < len(order) else len(lines)
        bodies[k] = "\n".join(lines[i:end]).strip()
    return found, bodies


def check_sections(text: str, doc_type: str):
    found, _ = section_bodies(text)
    idx = 3 if doc_type == "synopsis" else 4
    rows = []
    for s in SECTIONS:
        if s[idx]:
            rows.append({"key": s[0], "label": s[1], "present": s[0] in found})
    coverage = sum(r["present"] for r in rows) / max(len(rows), 1)
    return rows, coverage


def pick_excerpt(text: str, keys, frac=(0.0, 0.4), limit: int = EXCERPT_CHARS) -> str:
    _, bodies = section_bodies(text)
    parts = [bodies[k] for k in keys if k in bodies]
    if parts:
        per = max(limit // len(parts), 400)
        out = "\n\n".join(p[:per] for p in parts)
    else:
        n = len(text)
        out = text[int(n * frac[0]): int(n * frac[1])]
    return out[:limit]


def writing_stats(text: str) -> dict:
    words = word_count(text)
    sents = [s for s in re.split(r"(?<=[.!?])\s+", text) if len(s.split()) > 2]
    lens = [len(s.split()) for s in sents] or [0]
    passive = len(re.findall(r"\b(?:is|are|was|were|been|being|be)\s+\w+ed\b", text, re.I))
    cites = len(re.findall(r"\([A-Z][A-Za-z\-]+(?: et al\.)?[^()]{0,40}\d{4}[a-z]?\)", text)) + \
        len(re.findall(r"\[\d+(?:[,\u2013-]\s?\d+)*\]", text))
    _, bodies = section_bodies(text)
    refs = bodies.get("references", "")
    ref_entries = sum(1 for l in refs.split("\n")[1:] if re.search(r"\b(?:19|20)\d{2}\b", l))
    first_person = len(re.findall(r"\b(?:I|we|my|our)\b", text))
    return {
        "words": words,
        "avg_sentence_words": round(sum(lens) / len(lens), 1),
        "long_sentences_pct": round(100 * sum(l > 40 for l in lens) / len(lens), 1),
        "passive_per_1000_words": round(1000 * passive / max(words, 1), 1),
        "in_text_citations": cites,
        "reference_entries": ref_entries,
        "first_person_uses": first_person,
    }


# ------------------------------------------------------------------ LLM output
def parse_review(raw: str) -> dict:
    """Parse an agent's JSON review; never raises."""
    empty = {"score": None, "summary": "", "strengths": [], "issues": [], "parsed": False}
    if not raw:
        return empty
    s, e = raw.find("{"), raw.rfind("}")
    if s != -1 and e > s:
        try:
            d = json.loads(raw[s:e + 1])
            score = d.get("score")
            score = max(0.0, min(10.0, float(score))) if score is not None else None
            issues = [i for i in d.get("issues", []) if isinstance(i, dict)][:6]
            return {"score": score, "summary": str(d.get("summary", "")),
                    "strengths": [str(x) for x in d.get("strengths", [])][:4],
                    "issues": issues, "parsed": True}
        except (ValueError, TypeError):
            pass
    empty["summary"] = raw.strip()[:700]
    return empty


def compute_overall(scores: dict, coverage: float):
    """scores: key->float|None. Structure blends LLM score with Python coverage."""
    s = dict(scores)
    py_struct = coverage * 10
    s["structure"] = py_struct if s.get("structure") is None else 0.5 * py_struct + 0.5 * s["structure"]
    for k in WEIGHTS:
        if s.get(k) is None:
            s[k] = 5.0  # neutral if an agent output could not be parsed
    overall = sum(WEIGHTS[k] * s[k] for k in WEIGHTS)
    return round(overall, 1), {k: round(v, 1) for k, v in s.items()}


def decide(overall: float, coverage: float) -> str:
    if overall >= 8 and coverage >= 0.85:
        return "Approve"
    if overall >= 6.5:
        return "Approve with minor revisions"
    if overall >= 5:
        return "Major revision required"
    return "Not ready - substantial rework needed"


def input_hash(*parts) -> str:
    return hashlib.sha256("||".join(map(str, parts)).encode("utf-8")).hexdigest()


# ------------------------------------------------------------------- reporting
AGENT_LABELS = {"structure": "Document Structure Auditor", "problem": "Problem & Objectives Analyst",
                "method": "Methodology & Ethics Evaluator", "writing": "Academic Writing Reviewer"}


def build_report_md(meta: dict, scores: dict, reviews: dict, chair_text: str, sources: list) -> str:
    L = ["# ScholarPanel Review Report", "",
         f"- **Document type:** {meta['doc_type'].title()} ({meta['level']})",
         f"- **Words analysed:** {meta['words']}",
         f"- **Overall readiness score:** {meta['overall']} / 10",
         f"- **Recommended decision:** {meta['decision']}",
         f"- **Section coverage:** {round(meta['coverage'] * 100)}%", "",
         "> AI-generated advisory review. Final decisions rest with the supervisor and graduate committee.", "",
         "## Rubric Scores", ""]
    for k, lab in AGENT_LABELS.items():
        L.append(f"- {lab}: {scores.get(k)} / 10")
    L += ["", chair_text.strip(), "", "## Detailed Reviewer Findings"]
    for k, lab in AGENT_LABELS.items():
        r = reviews.get(k, {})
        L += ["", f"### {lab} ({scores.get(k)} / 10)", r.get("summary", "")]
        for i in r.get("issues", []):
            L.append(f"- **[{str(i.get('severity', '')).upper()}]** {i.get('issue', '')} "
                     f"-> *Fix:* {i.get('fix', '')}")
    if sources:
        L += ["", "## Guideline Sources Used"]
        for s in sources:
            L.append(f"- [{s['tag']}] {s['doc']} - {s['section']} (p. {s['page']})")
    return "\n".join(L)


def report_to_docx(md: str) -> bytes:
    d = Document()
    for line in md.split("\n"):
        t = line.rstrip()
        if t.startswith("# "):
            d.add_heading(t[2:], 0)
        elif t.startswith("## "):
            d.add_heading(t[3:], 1)
        elif t.startswith("### "):
            d.add_heading(t[4:], 2)
        elif re.match(r"^\s*[-*] ", t):
            d.add_paragraph(re.sub(r"[*_`>]", "", re.sub(r"^\s*[-*] ", "", t)), style="List Bullet")
        elif re.match(r"^\s*\d+[.)] ", t):
            d.add_paragraph(re.sub(r"[*_`>]", "", re.sub(r"^\s*\d+[.)] ", "", t)), style="List Number")
        elif t.strip():
            d.add_paragraph(re.sub(r"[*_`>]", "", t))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
